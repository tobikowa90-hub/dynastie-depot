"""Wegwerf-Skript Radikalkur Phase 1 Schritt 1: Inventar aller Alt-Regelquellen -> Abgleichliste.

Aufruf: python 02_Analysen/2026-10-17_inventar.py .   (im Repo-Root)
Schreibt 02_Analysen/2026-10-17_Abgleichliste.md und prueft danach Vollstaendigkeit
(jede Quell-ID genau einmal in der geschriebenen Datei; Ueberschriften-Zahl je Datei gegen Rohzaehlung).
"""
import re
import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(sys.argv[1])
OUT = ROOT / "02_Analysen" / "2026-10-17_Abgleichliste.md"
CONFIG = ROOT / "01_Skills" / "dynastie-depot" / "config.yaml"
XLSX = sorted((ROOT / "03_Tools").glob("*.xlsx"))
VAULT = ROOT / "07_Obsidian Vault"
WISS = VAULT / "Obsidian Mindmap/Investing Mastermind/wiki/synthesis/Wissenschaftliche-Fundierung-DEFCON.md"

# Regelartige Zeile: Schwelle, Vergleich, Punkte/Prozent/Tage/Mio, Malus/Bonus, DEFCON-Stufe
RULE = re.compile(r"(≥|≤|>=|<=|[<>] ?[0-9]|[0-9] ?(%|pp|Pt\.?|Punkte|Tage|Mio)|Schwelle|Malus|Bonus|DEFCON [1-4])")

# Markdown-Quellen: (Praefix, Pfad, Titel). Ueberschriften H2-H4 ausserhalb Code-Fences.
MD_SOURCES = [
    ("I", ROOT / "00_Core/INSTRUKTIONEN.md", "INSTRUKTIONEN.md"),
    ("K", ROOT / "00_Core/KONTEXT.md", "KONTEXT.md"),
    ("CM", ROOT / "00_Core/CORE-MEMORY.md", "CORE-MEMORY.md (→ Archiv E14)"),
    ("F", ROOT / "00_Core/Faktortabelle.md", "Faktortabelle.md (→ Archiv E14)"),
    ("P", ROOT / "00_Core/PORTFOLIO.md", "PORTFOLIO.md (→ Archiv E14)"),
    ("S", ROOT / "01_Skills/dynastie-depot/SKILL.md", "dynastie-depot SKILL.md"),
    ("B", ROOT / "01_Skills/dynastie-depot/Beispiele.md", "dynastie-depot Beispiele.md (Kalibrierungsanker)"),
    ("SR", ROOT / "01_Skills/dynastie-depot/sources.md", "dynastie-depot sources.md"),
    ("M", ROOT / "01_Skills/dynastie-depot/manifest.md", "dynastie-depot manifest.md"),
    ("QS", ROOT / "01_Skills/quick-screener/SKILL.md", "quick-screener SKILL.md"),
    ("QT", ROOT / "01_Skills/quick-screener/references/thresholds.md", "quick-screener references/thresholds.md"),
    ("II", ROOT / "01_Skills/insider-intelligence/SKILL.md", "insider-intelligence SKILL.md"),
    ("NU", ROOT / "01_Skills/non-us-fundamentals/SKILL.md", "non-us-fundamentals SKILL.md"),
    ("W", WISS, "Vault Wissenschaftliche-Fundierung-DEFCON.md (Überschriften)"),
    ("DS", ROOT / "03_Tools/depot-architecture-spec.md", "depot-architecture-spec.md (→ Archiv E14)"),
]
OVERRIDES = ROOT / "03_Tools/earnings_schedule_overrides.yaml"
SPEC = ROOT / "03_Tools/depot-architecture-spec.md"
# Code der aktiven Skills/Tools: alle Modul-Konstanten in GROSSBUCHSTABEN mit Literalwert (Codex-Review 18.10.)
CODE_ROOTS = [ROOT / "01_Skills/dynastie-depot", ROOT / "01_Skills/insider-intelligence",
              ROOT / "01_Skills/non-us-fundamentals", ROOT / "01_Skills/quick-screener", ROOT / "03_Tools"]


def cell(s, n=110):
    s = re.sub(r"\s+", " ", str(s)).strip().replace("|", "\\|")
    return s if len(s) <= n else s[: n - 1] + "…"


def rel(p):
    return p.relative_to(ROOT).as_posix()


def headings(path, prefix):
    """Ueberschriften H2-H4; Bereich = bis zur naechsten gleich-/hoeherrangigen,
    Treffer = Regelzeilen im *eigenen* Text bis zur naechsten Ueberschrift beliebiger Ebene."""
    lines = path.read_text(encoding="utf-8").splitlines()
    fence, heads = False, []
    for i, line in enumerate(lines, 1):
        if line.lstrip().startswith("```"):
            fence = not fence
            continue
        if not fence:
            m = re.match(r"^(#{2,4})\s+(.*)", line)
            if m:
                heads.append((i, len(m.group(1)), m.group(2)))
    if not heads:  # Klartext ohne Markdown-Struktur -> jede Regelzeile ist eine Einheit
        hits = [(i, l) for i, l in enumerate(lines, 1) if RULE.search(l)]
        return [(f"{prefix}-{n + 1:03d}", f"Z. {i} (Klartext)", cell(l, 120), "1") for n, (i, l) in enumerate(hits)]
    rows = []
    for n, (ln, lvl, title) in enumerate(heads):
        end = len(lines)
        for ln2, lvl2, _ in heads[n + 1:]:
            if lvl2 <= lvl:
                end = ln2 - 1
                break
        own_end = heads[n + 1][0] - 1 if n + 1 < len(heads) else len(lines)
        hits = sum(1 for l in lines[ln:own_end] if RULE.search(l))
        rows.append((f"{prefix}-{n + 1:03d}", f"Z. {ln}–{end}", "#" * lvl + " " + cell(title, 100), str(hits)))
    return rows


def raw_heading_count(path):
    """(alle H2-H4-Zeilen roh, davon innerhalb Code-Fences) — unabhaengig von headings() gezaehlt."""
    total = fenced = 0
    fence = False
    for l in path.read_text(encoding="utf-8").splitlines():
        if l.lstrip().startswith("```"):
            fence = not fence
        elif re.match(r"^#{2,4}\s", l):
            total += 1
            fenced += fence
    return total, fenced


def code_constants():
    import ast
    rows, files = [], []
    for r in CODE_ROOTS:
        for p in sorted(r.rglob("*.py")):
            if any(x in p.parts for x in ("tests", "test", "__pycache__", "fixtures")) or p.name.startswith("test_"):
                continue
            try:
                tree = ast.parse(p.read_text(encoding="utf-8"))
            except SyntaxError as e:
                sys.exit(f"FAIL: {rel(p)} nicht parsebar: {e}")
            hit = False
            for node in tree.body:
                tg = node.targets if isinstance(node, ast.Assign) else (
                    [node.target] if isinstance(node, ast.AnnAssign) and node.value else [])
                for x in tg:
                    if isinstance(x, ast.Name) and re.fullmatch(r"[A-Z][A-Z0-9_]+", x.id):
                        try:
                            val = ast.literal_eval(node.value)
                        except (ValueError, SyntaxError, TypeError):
                            continue
                        rows.append((f"PY-{len(rows) + 1:03d}", f"`{rel(p)}:{node.lineno}`", f"`{x.id}` = {cell(repr(val), 120)}", "–"))
                        hit = True
            if hit:
                files.append(rel(p))
    return rows, files


def fenced_blocks(path, prefix):
    lines = path.read_text(encoding="utf-8").splitlines()
    rows, start = [], None
    for i, line in enumerate(lines, 1):
        if line.lstrip().startswith("```"):
            if start is None:
                start = i
            else:
                body = [l for l in lines[start:i - 1] if l.strip()]
                rows.append((f"{prefix}-{len(rows) + 1:03d}", f"Z. {start}–{i}",
                             cell(lines[start - 1].strip() + " " + " · ".join(body[:6]), 160),
                             str(sum(1 for l in body if RULE.search(l)))))
                start = None
    return rows


def config_leaves(src=None, prefix="C"):
    src = src or CONFIG
    data = yaml.safe_load(src.read_text(encoding="utf-8"))
    acc = OrderedDict()

    def walk(x, path, owner):
        if isinstance(x, dict):
            for k, v in x.items():
                walk(v, f"{path}.{k}" if path else str(k), owner)
        elif isinstance(x, list) and any(isinstance(i, (dict, list)) for i in x):
            for v in x:
                tag = v.get("ticker") if isinstance(v, dict) else None
                walk(v, f"{path}[*]", tag or owner)
        else:
            acc.setdefault(path, []).append((owner, x))

    walk(data, "", None)
    rows = []
    for n, (path, vals) in enumerate(acc.items()):
        if len(vals) == 1:
            desc = cell(vals[0][1], 90)
        else:
            distinct = list(OrderedDict.fromkeys(cell(v, 30) for _, v in vals))
            desc = f"{len(vals)}× · {len(distinct)} Werte: " + " / ".join(distinct[:4]) + (" / …" if len(distinct) > 4 else "")
        rows.append((f"{prefix}-{n + 1:03d}", f"`{path}`", desc, "–"))
    return rows


def config_comments(src=None, prefix="CK"):
    lines = (src or CONFIG).read_text(encoding="utf-8").splitlines()
    blocks, cur = [], []
    for i, line in enumerate(lines, 1):
        s = line.strip()
        body = s.lstrip("#").strip() if s.startswith("#") else None
        if body and not re.fullmatch(r"[─\-=]+", body):
            cur.append((i, body))
            continue
        if cur:
            blocks.append(cur)
            cur = []
    if cur:
        blocks.append(cur)
    return [(f"{prefix}-{n + 1:03d}", f"Z. {b[0][0]}–{b[-1][0]}", cell(" ".join(t for _, t in b), 160),
             str(sum(1 for _, t in b if RULE.search(t)))) for n, b in enumerate(blocks)]


def xlsx_items():
    params, cfs, forms = [], [], OrderedDict()
    for f in XLSX:
        wb = openpyxl.load_workbook(f)
        for ws in wb:
            if ws.title == "Parameter & Regeln":
                for row in ws.iter_rows():
                    vals = [c.value for c in row]
                    if any(v is not None for v in vals):
                        params.append((f"{f.stem}!{ws.title}!{row[0].row}",
                                       " · ".join(cell(v, 60) for v in vals if v is not None)))
            for rng in ws.conditional_formatting:
                for r in rng.rules:
                    spec = r.formula or ([r.type] if r.type in ("dataBar", "colorScale", "iconSet") else [])
                    cfs.append((f"{f.stem}!{ws.title}!{rng.sqref}", f"{r.type}/{r.operator or '-'}: " + " ; ".join(map(str, spec))))
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if isinstance(v, str) and v.startswith("=") and re.search(r"[<>]=?|IF\(", v):
                        key = (f.stem, ws.title, re.sub(r"(?<=[A-Z$])\d+", "#", v))
                        forms.setdefault(key, []).append(c.coordinate)
    rows = [(f"X-{n + 1:03d}", f"`{loc}`", cell(d, 130), "–") for n, (loc, d) in enumerate(params)]
    rows += [(f"XF-{n + 1:03d}", f"`{loc}`", cell(d, 130), "–") for n, (loc, d) in enumerate(cfs)]
    rows += [(f"XL-{n + 1:03d}", f"`{k[0]}!{k[1]}!{cs[0]}`" + (f" (+{len(cs) - 1})" if len(cs) > 1 else ""), cell(k[2], 140), "–")
             for n, (k, cs) in enumerate(forms.items())]
    return rows


def vault_rules():
    rows = []
    files = sorted(p for p in VAULT.rglob("*.md")
                   if re.search(r"(?m)^operative_regel:", p.read_text(encoding="utf-8", errors="replace")))
    for n, p in enumerate(files):
        m = re.search(r'(?m)^operative_regel:\s*"?(.*?)"?\s*$', p.read_text(encoding="utf-8", errors="replace"))
        rows.append((f"V-{n + 1:03d}", f"`{p.relative_to(VAULT).as_posix()}`", cell(m.group(1), 160), "–"))
    return rows


SOURCES = [(f"{title} — Überschriften H2–H4", rel(path), headings(path, pre)) for pre, path, title in MD_SOURCES]
SOURCES += [
    ("config.yaml — Blätter (Listen als `[*]` zusammengefasst)", rel(CONFIG), config_leaves()),
    ("config.yaml — Kommentarblöcke", rel(CONFIG), config_comments()),
    ("xlsx — Parameterblatt / bedingte Formate / Vergleichsformeln", ", ".join(rel(f) for f in XLSX), xlsx_items()),
    ("Vault — `operative_regel`-Frontmatter", "07_Obsidian Vault/**/*.md", vault_rules()),
    ("depot-architecture-spec.md — Code-/YAML-Blöcke", rel(SPEC), fenced_blocks(SPEC, "DSY")),
    ("earnings_schedule_overrides.yaml — Blätter (ohne Regex-Filter)", rel(OVERRIDES), config_leaves(OVERRIDES, "EO")),
    ("earnings_schedule_overrides.yaml — Kommentarblöcke (ohne Regex-Filter)", rel(OVERRIDES), config_comments(OVERRIDES, "EOK")),
]
PY_ROWS, PY_FILES = code_constants()
SOURCES.append(("Code — Modul-Konstanten (GROSSBUCHSTABEN, Literalwert) aktiver Skills + 03_Tools", f"{len(PY_FILES)} .py-Dateien", PY_ROWS))

# ---------- Suchraum-Komplement: kompletter Dateibestand, auch 0 Treffer ----------
included = {rel(p) for _, p, _ in MD_SOURCES} | {rel(CONFIG), rel(OVERRIDES)} | {rel(f) for f in XLSX} | set(PY_FILES)
cands = [ROOT / "CLAUDE.md"]
for d in ("00_Core", "01_Skills", "03_Tools", "04_Templates", ".claude"):
    for p in (ROOT / d).rglob("*"):
        if p.is_file() and p.suffix in (".md", ".yaml", ".yml", ".py", ".json", ".toml") and "_extern" not in p.parts \
                and "__pycache__" not in p.parts and "workspace" not in "/".join(p.parts):
            cands.append(p)
seen, uniq = set(), []
for p in cands:  # .claude/skills/* sind Junctions auf 01_Skills -> ueber aufgeloesten Pfad entdoppeln
    r = p.resolve()
    if r not in seen:
        seen.add(r)
        uniq.append(p if r == p.absolute() else r)
cands = uniq
compl = []
for p in cands:
    txt = p.read_text(encoding="utf-8", errors="replace").splitlines()
    compl.append((rel(p), len(txt), sum(1 for l in txt if RULE.search(l)), rel(p) in included))
compl.sort(key=lambda r: (-r[2], r[0]))

out = ["# Abgleichliste Alt-Regelquellen — Radikalkur Phase 1", "",
       "> Erzeugt 2026-10-18 per `02_Analysen/2026-10-17_inventar.py` (Neuerzeugung: `python 02_Analysen/2026-10-17_inventar.py .` im Repo-Root; Codex-Scope-Review 18.10. eingearbeitet). Schritt 1 = Inventar; Spalten **Verdikt / E-Bezug / Skill wörtlich?** füllt Schritt 2 (Verfahren E18).",
       "> Verdikt-Werte: `übernehmen` · `streichen` · `Zustand→Broker` · `Urteil→Ledger` · `Methode→Skill` · `offen → Owner`.",
       "> **Treffer** = regelartige Zeilen im eigenen Abschnittstext (Muster: ≥ ≤ < > · % pp Pt. Punkte Tage Mio · Schwelle Malus Bonus · DEFCON 1–4). Abschnitte mit vielen Treffern werden in Schritt 2 in Einzelregeln zerlegt — „Fehlende = 0“ gilt auf Abschnittsebene, nicht auf Regelebene.", "",
       "## Suchraum", "",
       "| Quelle | Einheit | Anzahl |", "|---|---|---|"]
for title, path, rows in SOURCES:
    out.append(f"| `{cell(path, 200)}` | {title} | {len(rows)} |")
total = sum(len(r) for _, _, r in SOURCES)
out += [f"| **Σ** | | **{total}** |", "",
        "**Bewusst nicht im Suchraum (Owner 18.10.):** H1-Dateititel · `CLAUDE.md`, `APPLIED-LEARNING.md`, `STATE.md` (Ablaufregeln → E15/E16) · Memories (E16) · xlsx-Datenvalidierungen (Dropdowns) und US-Exposure-Blatt (reine Verweise) · Vault-Fließtext außer `operative_regel` + Wiss-Fundierung-Überschriften. Vollständige Gegenliste → letzter Abschnitt.",
        ""]
for title, _, rows in SOURCES:
    out += [f"## {title}", "", "| ID | Fundstelle | Inhalt | Treffer | Verdikt | E-Bezug | Skill wörtlich? |",
            "|---|---|---|---|---|---|---|"]
    out += [f"| {i} | {loc} | {desc} | {hits} | | | |" for i, loc, desc, hits in rows]
    out.append("")
out += ["## Suchraum-Komplement (Dateibestand 00_Core · 01_Skills ohne `_extern` · 03_Tools · 04_Templates · .claude · CLAUDE.md)", "",
        "> Mechanisch erzeugt: **jede** .md/.yaml/.yml/.py/.json/.toml-Datei, auch mit 0 Treffern (Regex = Priorisierung, kein Vollständigkeitsnachweis — Textregeln wie „KEIN“/„Tag +1“ trifft er nicht). `✓` = oben inventarisiert. Nicht inventarisierte Dateien sind in Schritt 2 zu begründen (Archiv, Ablaufregel, Infrastruktur).", "",
        "| Datei | Zeilen | Treffer | im Suchraum |", "|---|---|---|---|"]
out += [f"| `{p}` | {n} | {h} | {'✓' if inc else ''} |" for p, n, h, inc in compl]
out.append("")
OUT.write_text("\n".join(out) + "\n", encoding="utf-8", newline="\n")

# ---------- Pruefung ----------
leer = [title for title, _, rows in SOURCES if not rows]
if leer:
    sys.exit(f"FAIL: Quelle ohne Einheiten (stiller Null-Treffer): {leer}")
expected = [r[0] for _, _, rows in SOURCES for r in rows]
written = re.findall(r"(?m)^\| ((?:I|K|CM|F|P|S|B|SR|M|QS|QT|II|NU|W|DS|DSY|C|CK|X|XF|XL|V|EO|EOK|PY)-\d{3}) \|", OUT.read_text(encoding="utf-8"))
missing, extra = set(expected) - set(written), set(written) - set(expected)
dupes = {w for w in written if written.count(w) > 1}
print(f"Quell-Einheiten: {len(expected)} | geschrieben: {len(written)} | Fehlende: {len(missing)} | Doppelte: {len(dupes)} | Fremde: {len(extra)}")
for (pre, path, _), (title, _, rows) in zip(MD_SOURCES, SOURCES):
    raw, fenced = raw_heading_count(path)
    print(f"  {len(rows):4d} (roh {raw:4d}, in Fences {fenced}{' ABWEICHUNG' if raw - fenced != len(rows) else ''})  {title}")
for title, _, rows in SOURCES[len(MD_SOURCES):]:
    print(f"  {len(rows):4d}  {title}")
nicht = [(p, h) for p, _, h, inc in compl if h > 0 and not inc]
print(f"Komplement: {len(nicht)} Dateien mit Treffern nicht inventarisiert; Top 15:")
for p, h in nicht[:15]:
    print(f"  {h:4d}  {p}")
